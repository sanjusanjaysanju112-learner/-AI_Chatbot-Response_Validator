"""
Turns the log-triage results into a 3-slide .pptx deck.

Uses python-pptx (pure Python, pip-installable, no Node/npm dependency) -
deliberately chosen over JS-based generators since you may not have full
control over the interview machine's setup on the day.

The AI does the *writing* (via generate_slide_copy, optional) - this module
does the *building* (turning that text into an actual PowerPoint file).
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from io import BytesIO

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x2E, 0x6E, 0xE8)


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # fully blank layout


def _add_title(slide, text, top=0.5):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(9.0), Inches(1.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = NAVY
    return box


def _add_bullets(slide, items, top=1.7, width=8.8, font_size=18):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(width), Inches(4.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = GRAY
        p.space_after = Pt(10)
    return box


def build_deck(problem_statement: str, architecture_steps: list, severity_counts: dict,
               impact_points: list) -> bytes:
    """
    problem_statement: one paragraph describing the problem
    architecture_steps: ordered list of strings, e.g.
        ["Logs pulled from S3 / CloudWatch", "LLM classifies failure type,
         root cause, severity", "Results shown in live dashboard"]
    severity_counts: dict like {"Critical": 2, "High": 4, "Medium": 3, "Low": 1}
    impact_points: list of strings for the results/impact slide
    """
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    # --- Slide 1: Problem ---
    s1 = _blank_slide(prs)
    s1.background.fill.solid()
    s1.background.fill.fore_color.rgb = WHITE
    _add_title(s1, "AI-Powered Defect / Log Triage", top=1.6)
    sub = s1.shapes.add_textbox(Inches(0.6), Inches(2.6), Inches(8.8), Inches(2.0))
    tf = sub.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = problem_statement
    p.font.size = Pt(18)
    p.font.color.rgb = GRAY

    # --- Slide 2: Solution / Architecture ---
    s2 = _blank_slide(prs)
    s2.background.fill.solid()
    s2.background.fill.fore_color.rgb = WHITE
    _add_title(s2, "Solution & Architecture")
    _add_bullets(s2, architecture_steps, top=1.7)

    # --- Slide 3: Results & Impact ---
    s3 = _blank_slide(prs)
    s3.background.fill.solid()
    s3.background.fill.fore_color.rgb = WHITE
    _add_title(s3, "Results & Impact")

    # Severity counts as simple metric row
    order = ["Critical", "High", "Medium", "Low"]
    n = len(order)
    box_w = 2.0
    gap = 0.25
    start_x = (10 - (n * box_w + (n - 1) * gap)) / 2
    for i, sev in enumerate(order):
        x = start_x + i * (box_w + gap)
        box = s3.shapes.add_textbox(Inches(x), Inches(1.8), Inches(box_w), Inches(1.2))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.text = str(severity_counts.get(sev, 0))
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = ACCENT if sev in ("Critical", "High") else GRAY
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = sev
        p2.font.size = Pt(14)
        p2.font.color.rgb = GRAY
        p2.alignment = PP_ALIGN.CENTER

    _add_bullets(s3, impact_points, top=3.3, font_size=16)

    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    # Quick smoke test with sample data
    deck_bytes = build_deck(
        problem_statement=("Manual triage of failure logs across payment services is slow "
                            "and inconsistent between engineers, delaying incident response."),
        architecture_steps=[
            "Logs ingested from S3 / CloudWatch (fetch_logs_from_cloud)",
            "Each entry classified by an LLM: failure type, root cause, severity",
            "Rule-based fallback mode keeps the tool working if the AI API is unavailable",
            "Results shown in a live dashboard with severity counts and CSV export",
        ],
        severity_counts={"Critical": 3, "High": 4, "Medium": 2, "Low": 1},
        impact_points=[
            "First-pass severity and root-cause hypothesis in seconds instead of manual log reading",
            "Engineers triage by exception, focusing first on Critical/High items",
            "Consistent classification across engineers and shifts",
        ],
    )
    with open("demo_deck.pptx", "wb") as f:
        f.write(deck_bytes)
    print(f"Wrote demo_deck.pptx ({len(deck_bytes)} bytes)")
